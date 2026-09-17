import json
import sys
import os
from openpyxl import Workbook, load_workbook
from datetime import datetime
from openpyxl.styles import Alignment

def resource_path(relative_path):
    if hasattr(sys, '_MEIPASS'):
        base_path = sys._MEIPASS
    else:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)

def get_app_data_dir():
    if sys.platform == "win32":
        app_data = os.getenv("APPDATA")
        dir_name = "TestApp"
    elif sys.platform == "darwin":
        app_data = os.path.expanduser("~/Library/Application Support")
        dir_name = "TestApp"
    else:
        app_data = os.path.expanduser("~/.local/share")
        dir_name = "testapp"
        
    path = os.path.join(app_data, dir_name)
    if not os.path.exists(path):
        os.makedirs(path)
    return path

OPTIONS = [
    {"text": "Да", "score": 2},
    {"text": "Иногда", "score": 1},
    {"text": "Нет", "score": 0}
]

SCALE_NAMES = {
    "attitude_to_victory": "Отношение к победе",
    "reaction_to_defeat": "Реакция на поражение",
    "reflection": "Рефлексия и анализ",
    "emotional_stability": "Эмоциональная устойчивость",
    "social_maturity": "Социальная зрелость"
}

class UserData:
    def __init__(self):
        self.name = ""
        self.age = ""

user_data = UserData()

def load_questions():
    json_path = resource_path('questions.json')
    
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        print(f"Ошибка загрузки: {e}")
        return []

def get_options():
    return OPTIONS

def rec_attitude_to_victory(percentage: float) -> str:
    if percentage >= 75:
        return "Высокий уровень."
    elif percentage >= 37.5:
        return "Средний уровень."
    else:
        return "Низкий уровень."

def rec_reaction_to_defeat(percentage: float) -> str:
    if percentage >= 75:
        return "Высокий уровень."
    elif percentage >= 37.5:
        return "Средний уровень."
    else:
        return "Низкий уровень."

def rec_reflection(percentage: float) -> str:
    if percentage >= 75:
        return "Высокий уровень."
    elif percentage >= 37.5:
        return "Средний уровень."
    else:
        return "Низкий уровень."

def rec_emotional_stability(percentage: float) -> str:
    if percentage >= 75:
        return "Высокий уровень."
    elif percentage >= 37.5:
        return "Средний уровень."
    else:
        return "Низкий уровень."

def rec_social_maturity(percentage: float) -> str:
    if percentage >= 75:
        return "Высокий уровень."
    elif percentage >= 37.5:
        return "Средний уровень."
    else:
        return "Низкий уровень."

RECOMMENDATION_FUNCTIONS = {
    "attitude_to_victory": rec_attitude_to_victory,
    "reaction_to_defeat": rec_reaction_to_defeat,
    "reflection": rec_reflection,
    "emotional_stability": rec_emotional_stability,
    "social_maturity": rec_social_maturity
}

def calculate_profile(questions: list, user_answers_indices: list) -> dict:
    scales = {}
    total_score = 0
    total_max = 0
    
    for q in questions:
        scale_key = q.get("scale", "unknown")
        if scale_key not in scales:
            scales[scale_key] = {
                "name": SCALE_NAMES.get(scale_key, scale_key),
                "score": 0,
                "max": 0,
            }
        
        idx = questions.index(q)
        if idx < len(user_answers_indices):
            ans_idx = user_answers_indices[idx]
            option = OPTIONS[ans_idx]
            
            scales[scale_key]["score"] += option["score"]
            scales[scale_key]["max"] += 2
            
            total_score += option["score"]
            total_max += 2

    results_list = []
    for key, data in scales.items():
        percentage = (data["score"] / data["max"] * 100) if data["max"] > 0 else 0

        rec_func = RECOMMENDATION_FUNCTIONS.get(key)
        if rec_func:
            recommendation = rec_func(percentage)
        else:
            recommendation = "Нет данных для рекомендации."
        
        results_list.append({
            "scale_name": data["name"],
            "score": data["score"],
            "max": data["max"],
            "percentage": percentage,
            "recommendation": recommendation
        })
        
    order = ["attitude_to_victory", "reaction_to_defeat", "reflection", "emotional_stability", "social_maturity"]
    results_list.sort(key=lambda x: order.index(x["scale_name"]) if x["scale_name"] in order else 99)

    total_percentage = (total_score / total_max * 100) if total_max > 0 else 0
    
    return {
        "scales": results_list,
        "total_score": total_score,
        "total_max": total_max,
        "total_percentage": total_percentage
    }

def save_to_excel(user_name, user_age, profile_data):
    filename = "results.xlsx"
    folder_path = get_app_data_dir()
    filepath = os.path.join(folder_path, filename)
    
    headers = [
        "Дата","ФИО", "Возраст", 
        "Победа", "Поражение", "Рефлексия", "Устойчивость", "Зрелость", "Всего баллов"
    ]

    center_alignment = Alignment(horizontal="center", vertical="center")
    
    if os.path.exists(filepath):
        wb = load_workbook(filepath)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"
        ws.append(headers)

        for cell in ws[1]:
            cell.alignment = center_alignment

    scale_scores = [0, 0, 0, 0, 0]
    
    for item in profile_data['scales']:
        name = item['scale_name']
        score = item['score']
        
        if "победе" in name: 
            scale_scores[0] = score
        elif "поражение" in name: 
            scale_scores[1] = score
        elif "Рефлексия" in name: 
            scale_scores[2] = score
        elif "устойчивость" in name: 
            scale_scores[3] = score
        elif "зрелость" in name: 
            scale_scores[4] = score

    total_sum = sum(scale_scores)

    row_data = [
        datetime.now().strftime("%d.%m.%y"),
        user_name,
        user_age,
        *scale_scores,
        total_sum
    ]

    ws.append(row_data)
    
    last_row_idx = ws.max_row
    for cell in ws[last_row_idx]:
        cell.alignment = center_alignment

    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(filepath)
    return filepath

def delete_last_record():
    filename = "results.xlsx"
    folder_path = get_app_data_dir()
    filepath = os.path.join(folder_path, filename)
    
    if not os.path.exists(filepath):
        return False, "Файл с результатами не найден."
        
    try:
        wb = load_workbook(filepath)
        ws = wb.active
        
        if ws.max_row <= 1:
            return False, "В файле нет записей для удаления."
            
        ws.delete_rows(ws.max_row, 1)
        wb.save(filepath)
        return True, "Последняя запись успешно удалена."
        
    except Exception as e:
        return False, f"Ошибка при удалении: {str(e)}"